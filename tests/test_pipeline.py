from pathlib import Path
from unittest.mock import patch

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from PIL import Image
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.database import get_db
from app.main import app
from app.models.document import Document, DocumentStatus
from app.models.group import Base, Group
from app.services.converter import ConversionResult, ConversionService
from app.services.pipeline import PipelineService
from app.services.preprocessing import PreprocessingService
from app.services.storage import StorageService


# --- Fixtures ---


@pytest.fixture
def synthetic_image() -> Image.Image:
    """Create a synthetic RGB image for testing."""
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    return img


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """Generate a minimal valid PDF file for testing."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello World from PDF")
    pdf_path = tmp_path / "sample.pdf"
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Create a minimal DOCX file for testing."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("This is test paragraph content.")
    docx_path = tmp_path / "sample.docx"
    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def sample_image(tmp_path: Path) -> Path:
    """Create a sample image file for testing."""
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    img_path = tmp_path / "sample.png"
    img.save(str(img_path))
    return img_path


@pytest_asyncio.fixture
async def pipeline_db_session(tmp_path: Path):
    """Create a file-based DB session for pipeline tests."""
    db_path = tmp_path / "pipeline_test.db"
    db_url = f"sqlite+aiosqlite:///{db_path}"
    engine = create_async_engine(db_url, echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    session_factory = async_sessionmaker(
        engine, class_=AsyncSession, expire_on_commit=False
    )
    async with session_factory() as session:
        yield session

    await engine.dispose()


@pytest_asyncio.fixture
async def pipeline_client(
    pipeline_db_session: AsyncSession, tmp_path: Path
) -> AsyncClient:
    """Create a test client with pipeline support."""
    from collections.abc import AsyncGenerator

    db_url = str(pipeline_db_session.get_bind().url)

    async def override_get_db() -> AsyncGenerator[AsyncSession, None]:
        try:
            yield pipeline_db_session
            await pipeline_db_session.commit()
        except Exception:
            await pipeline_db_session.rollback()
            raise

    app.dependency_overrides[get_db] = override_get_db

    import app.api.documents as doc_module

    doc_module.storage_service = StorageService(base_path=str(tmp_path / "storage"))
    doc_module.pipeline_db_url = db_url
    doc_module.pipeline_service = PipelineService(
        storage_service=doc_module.storage_service
    )

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac

    app.dependency_overrides.clear()


# --- Preprocessing Tests ---


class TestPreprocessing:
    def test_enhance_image(self, synthetic_image: Image.Image):
        service = PreprocessingService()
        result = service.enhance_image(synthetic_image)
        assert isinstance(result, Image.Image)
        assert result.size == synthetic_image.size

    def test_denoise_image(self, synthetic_image: Image.Image):
        service = PreprocessingService()
        result = service.denoise_image(synthetic_image)
        assert isinstance(result, Image.Image)
        assert result.size == synthetic_image.size

    def test_deskew_image_without_tesseract(self, synthetic_image: Image.Image):
        """Test that deskew handles pytesseract failure gracefully."""
        service = PreprocessingService()
        with patch(
            "pytesseract.image_to_osd", side_effect=Exception("tesseract not found")
        ):
            result = service.deskew_image(synthetic_image)
        assert isinstance(result, Image.Image)
        assert result.size == synthetic_image.size

    def test_deskew_image_graceful_fallback(self, synthetic_image: Image.Image):
        """Test that deskew gracefully handles pytesseract failure."""
        service = PreprocessingService()
        with patch(
            "pytesseract.image_to_osd", side_effect=Exception("tesseract not found")
        ):
            result = service.deskew_image(synthetic_image)
        assert isinstance(result, Image.Image)
        assert result.size == synthetic_image.size

    def test_preprocess_pipeline(self, synthetic_image: Image.Image):
        """Test the full preprocessing pipeline."""
        service = PreprocessingService()
        with patch(
            "pytesseract.image_to_osd", side_effect=Exception("not available")
        ):
            result = service.preprocess_pipeline(synthetic_image)
        assert isinstance(result, Image.Image)


# --- Converter Tests ---


class TestConversionService:
    def test_convert_pdf(self, sample_pdf: Path):
        """Test PDF conversion extracts text."""
        service = ConversionService()
        result = service.convert_pdf(sample_pdf)
        assert isinstance(result, ConversionResult)
        assert "Hello World from PDF" in result.markdown_content

    def test_convert_docx(self, sample_docx: Path):
        """Test DOCX conversion extracts text and headings."""
        service = ConversionService()
        result = service.convert_docx(sample_docx)
        assert isinstance(result, ConversionResult)
        assert "Test Heading" in result.markdown_content
        assert "test paragraph content" in result.markdown_content
        # Verify heading is formatted as markdown
        assert "# Test Heading" in result.markdown_content

    @patch("pytesseract.image_to_string", return_value="OCR extracted text")
    def test_convert_image_with_tesseract(self, mock_ocr, sample_image: Path):
        """Test image conversion with mocked tesseract."""
        service = ConversionService()
        with patch(
            "pytesseract.image_to_osd", side_effect=Exception("not available")
        ):
            result = service.convert_image(sample_image)
        assert isinstance(result, ConversionResult)
        assert "OCR extracted text" in result.markdown_content
        mock_ocr.assert_called_once()

    @patch(
        "pytesseract.image_to_string",
        side_effect=Exception("tesseract not installed"),
    )
    def test_convert_image_without_tesseract(self, mock_ocr, sample_image: Path):
        """Test image conversion gracefully handles missing tesseract."""
        service = ConversionService()
        with patch(
            "pytesseract.image_to_osd", side_effect=Exception("not available")
        ):
            result = service.convert_image(sample_image)
        assert isinstance(result, ConversionResult)
        assert "warning" in result.metadata

    def test_convert_pdf_no_tesseract_fallback(self, sample_pdf: Path):
        """Test PDF conversion still works when tesseract is unavailable."""
        service = ConversionService()
        # pymupdf extracts text directly, no tesseract needed
        result = service.convert_pdf(sample_pdf)
        assert isinstance(result, ConversionResult)
        assert "Hello World from PDF" in result.markdown_content


# --- Pipeline Tests ---


class TestPipelineService:
    @pytest.mark.asyncio
    async def test_process_document_pdf(
        self, pipeline_db_session: AsyncSession, sample_pdf: Path, tmp_path: Path
    ):
        """Test pipeline processes a PDF document."""
        storage_path = tmp_path / "storage"
        storage_service = StorageService(base_path=str(storage_path))

        # Create group and document
        group = Group(name="Test Group")
        pipeline_db_session.add(group)
        await pipeline_db_session.flush()
        await pipeline_db_session.refresh(group)

        document = Document(
            group_id=group.id,
            original_filename="sample.pdf",
            storage_path=str(sample_pdf),
            file_type="application/pdf",
            file_size=sample_pdf.stat().st_size,
            status=DocumentStatus.processing,
        )
        pipeline_db_session.add(document)
        await pipeline_db_session.flush()
        await pipeline_db_session.refresh(document)

        pipeline = PipelineService(storage_service=storage_service)
        await pipeline.process_document(document.id, pipeline_db_session)

        await pipeline_db_session.refresh(document)
        assert document.status == DocumentStatus.processed
        assert document.markdown_path is not None
        assert Path(document.markdown_path).exists()
        content = Path(document.markdown_path).read_text()
        assert "Hello World from PDF" in content

    @pytest.mark.asyncio
    async def test_process_document_not_found(
        self, pipeline_db_session: AsyncSession, tmp_path: Path
    ):
        """Test pipeline handles missing document gracefully."""
        storage_service = StorageService(base_path=str(tmp_path / "storage"))
        pipeline = PipelineService(storage_service=storage_service)
        # Should not raise, just log error
        await pipeline.process_document(9999, pipeline_db_session)

    @pytest.mark.asyncio
    async def test_process_document_failure(
        self, pipeline_db_session: AsyncSession, tmp_path: Path
    ):
        """Test pipeline sets status to failed on error."""
        storage_path = tmp_path / "storage"
        storage_service = StorageService(base_path=str(storage_path))

        group = Group(name="Test Group")
        pipeline_db_session.add(group)
        await pipeline_db_session.flush()
        await pipeline_db_session.refresh(group)

        # Point to a non-existent file to trigger failure
        document = Document(
            group_id=group.id,
            original_filename="nonexistent.pdf",
            storage_path="/nonexistent/path/file.pdf",
            file_type="application/pdf",
            file_size=0,
            status=DocumentStatus.processing,
        )
        pipeline_db_session.add(document)
        await pipeline_db_session.flush()
        await pipeline_db_session.refresh(document)

        pipeline = PipelineService(storage_service=storage_service)
        await pipeline.process_document(document.id, pipeline_db_session)

        await pipeline_db_session.refresh(document)
        assert document.status == DocumentStatus.failed


# --- API Endpoint Tests ---


class TestStatusEndpoint:
    @pytest.mark.asyncio
    async def test_get_status(self, pipeline_client: AsyncClient, pipeline_db_session: AsyncSession):
        """Test GET /api/documents/{id}/status returns status."""
        # Create group first
        response = await pipeline_client.post(
            "/api/groups", json={"name": "Status Test Group"}
        )
        assert response.status_code == 201
        group_id = response.json()["id"]

        # Upload a document
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Status test content")
        pdf_bytes = doc.tobytes()
        doc.close()

        response = await pipeline_client.post(
            f"/api/groups/{group_id}/documents",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
        )
        assert response.status_code == 201
        doc_id = response.json()["id"]

        # Check status
        response = await pipeline_client.get(f"/api/documents/{doc_id}/status")
        assert response.status_code == 200
        data = response.json()
        assert data["id"] == doc_id
        assert data["status"] in ["processing", "processed"]

    @pytest.mark.asyncio
    async def test_get_status_not_found(self, pipeline_client: AsyncClient):
        """Test GET /api/documents/{id}/status returns 404 for missing doc."""
        response = await pipeline_client.get("/api/documents/99999/status")
        assert response.status_code == 404


class TestMarkdownEndpoint:
    @pytest.mark.asyncio
    async def test_get_markdown_not_processed(
        self, pipeline_client: AsyncClient, pipeline_db_session: AsyncSession
    ):
        """Test markdown endpoint returns 400 if document not processed."""
        response = await pipeline_client.post(
            "/api/groups", json={"name": "Markdown Test Group"}
        )
        assert response.status_code == 201
        group_id = response.json()["id"]

        # Upload a simple text file (not PDF) that won't be auto-processed into markdown
        response = await pipeline_client.post(
            f"/api/groups/{group_id}/documents",
            files={"file": ("test.txt", b"hello", "text/plain")},
        )
        assert response.status_code == 201
        doc_id = response.json()["id"]

        # Manually set status back to uploaded for this test
        doc = await pipeline_db_session.get(Document, doc_id)
        doc.status = DocumentStatus.uploaded
        await pipeline_db_session.commit()

        response = await pipeline_client.get(f"/api/documents/{doc_id}/markdown")
        assert response.status_code == 400

    @pytest.mark.asyncio
    async def test_get_markdown_success(
        self, pipeline_client: AsyncClient, pipeline_db_session: AsyncSession, tmp_path: Path
    ):
        """Test markdown endpoint returns content for processed document."""
        response = await pipeline_client.post(
            "/api/groups", json={"name": "Markdown Success Group"}
        )
        assert response.status_code == 201
        group_id = response.json()["id"]

        # Upload
        response = await pipeline_client.post(
            f"/api/groups/{group_id}/documents",
            files={"file": ("test.txt", b"hello", "text/plain")},
        )
        assert response.status_code == 201
        doc_id = response.json()["id"]

        # Manually set document as processed with a markdown file
        doc = await pipeline_db_session.get(Document, doc_id)
        md_path = tmp_path / "test_markdown.md"
        md_path.write_text("# Test Content\n\nHello World", encoding="utf-8")
        doc.status = DocumentStatus.processed
        doc.markdown_path = str(md_path)
        await pipeline_db_session.commit()

        response = await pipeline_client.get(f"/api/documents/{doc_id}/markdown")
        assert response.status_code == 200
        assert "Test Content" in response.text
        assert "Hello World" in response.text

    @pytest.mark.asyncio
    async def test_get_markdown_not_found(self, pipeline_client: AsyncClient):
        """Test markdown endpoint returns 404 for missing doc."""
        response = await pipeline_client.get("/api/documents/99999/markdown")
        assert response.status_code == 404
